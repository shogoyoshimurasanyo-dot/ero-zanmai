#!/usr/bin/env python3
"""
毎日 8:00 に Windows のタスクスケジューラから実行する想定のスクリプト。

トップページ（index.html）の動画グリッドを毎日入れ替える:
  1. Excel（シート「API」）から C 列＝タイトル、D 列＝品番(cid) を読む
  2. サーバから index.html とクリックログ（api/clicks.log）を FTP でダウンロード
  3. 直近 N 日のクリックが多い動画を上位 keep_top 件まで「残す」
  4. 残りの枠を Excel からランダムに選んだ動画で埋める（最近出したものは避ける）
  5. STAGE 区間を作り直して index.html を FTP でアップロード
     （クリック計測用の click-track.js / api/click.php も一緒にアップロード）

使い方:
  py zanmai_daily.py            本番実行
  py zanmai_daily.py --dry-run  アップロードせず preview.html だけ作る
  py zanmai_daily.py --force    今日すでに実行済みでももう一度実行する
"""
import argparse
import datetime as dt
import ftplib
import html
import io
import json
import posixpath
import random
import re
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
JST = dt.timezone(dt.timedelta(hours=9))
NOW = dt.datetime.now(JST)
TODAY = NOW.date()
TRACK_TAG = '<script src="/assets/js/click-track.js" defer></script>'
SITE_TAG = '<script src="/assets/js/site.js" defer></script>'  # 説明文の「続きを読む」


def log(msg):
    line = f"[{NOW:%Y-%m-%d %H:%M:%S}] {msg}"
    print(line)
    (HERE / "logs").mkdir(exist_ok=True)
    with open(HERE / "logs" / f"{NOW:%Y-%m}.log", "a", encoding="utf-8") as f:
        f.write(line + "\n")


def load_config():
    path = HERE / "config.json"
    if not path.exists():
        sys.exit("config.json がありません。config.example.json をコピーして作ってください。")
    return json.loads(path.read_text(encoding="utf-8"))


# ---------- Excel ----------

def find_excel(path_str):
    p = Path(path_str)
    if p.is_file():
        return p
    if p.is_dir():
        files = [f for f in p.iterdir()
                 if f.suffix.lower() in (".xlsx", ".xlsm") and not f.name.startswith("~$")]
        if len(files) == 1:
            return files[0]
        names = ", ".join(f.name for f in files) or "（なし）"
        sys.exit(f"{p} の Excel ファイルを 1 つに特定できません: {names}\n"
                 "config.json の excel_path にファイル名まで書いてください。")
    sys.exit(f"excel_path が見つかりません: {p}")


def col_index(letter):
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def read_rows(cfg):
    """[(key, title, value), ...] を返す。key は重複判定用（cid）。"""
    from openpyxl import load_workbook  # pip install openpyxl

    src = find_excel(cfg["excel_path"])
    ti, vi = col_index(cfg["title_column"]), col_index(cfg["value_column"])
    rows, seen = [], set()
    # Excel で開いたままでも読めるよう、一時フォルダにコピーしてから読む
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / src.name
        shutil.copy2(src, copy)
        wb = load_workbook(copy, read_only=True, data_only=True)
        if cfg["sheet"] not in wb.sheetnames:
            sys.exit(f"シート「{cfg['sheet']}」がありません。存在するシート: {wb.sheetnames}")
        for r in wb[cfg["sheet"]].iter_rows(min_row=cfg.get("start_row", 2),
                                            min_col=min(ti, vi), max_col=max(ti, vi),
                                            values_only=True):
            title = r[ti - min(ti, vi)]
            value = r[vi - min(ti, vi)]
            if value is None or str(value).strip() == "":
                continue
            value = str(value).strip()
            key = item_key(value)
            if key in seen:
                continue
            seen.add(key)
            rows.append((key, "" if title is None else str(title).strip(), value))
        wb.close()
    log(f"Excel 読み込み: {src.name} / {cfg['sheet']} C=タイトル, D=値 → {len(rows)} 件（重複除く）")
    return rows


# ---------- HTML ----------

CID_RE = re.compile(r"cid=([A-Za-z0-9_]+)")


def item_key(text):
    """同じ作品かどうかの判定キー。cid があれば cid、なければ値そのもの。"""
    m = CID_RE.search(text)
    if m:
        return m.group(1)
    return text.strip()


def render(title, value, cfg):
    if value.lstrip().startswith("<"):
        return value  # D 列が HTML ならそのまま使う
    return (cfg["item_template"]
            .replace("{value}", value)
            .replace("{title}", html.escape(title)))


def stage_bounds(page, cfg):
    """STAGE 区間の <ul ...> の直後と </ul> の位置を返す。"""
    start = page.find(cfg["start_marker"])
    end = page.find(cfg["end_marker"])
    if start == -1 or end == -1 or end < start:
        sys.exit(f"HTML に「{cfg['start_marker']}」「{cfg['end_marker']}」が見つかりません。中断します。")
    ul_open = page.find("<ul", start, end)
    close_ul = page.rfind("</ul>", start, end)
    if ul_open == -1 or close_ul == -1:
        sys.exit("STAGE 区間の中に <ul>〜</ul> が見つかりません。中断します。")
    inner_start = page.index(">", ul_open) + 1
    return inner_start, close_ul


def current_items(page, cfg):
    """今ページに出ている動画 {key: <li>…</li>} （表示順）"""
    a, b = stage_bounds(page, cfg)
    items = {}
    for m in re.finditer(r"<li\b.*?</li>", page[a:b], re.S):
        items.setdefault(item_key(m.group(0)), m.group(0))
    return items


def rebuild_page(page, items_html, cfg):
    a, b = stage_bounds(page, cfg)
    body = "\n" + "\n".join("      " + h.strip() for h in items_html) + "\n    "
    page = page[:a] + body + page[b:]
    if "/assets/js/site.js" not in page:
        page = page.replace("</body>", f"{SITE_TAG}\n</body>", 1)
    if cfg.get("tracking", True) and "click-track.js" not in page:
        page = page.replace("</body>", f"{TRACK_TAG}\n</body>", 1)
    return page


# ---------- クリック集計・選定 ----------

def count_clicks(log_text, days):
    since = TODAY - dt.timedelta(days=days - 1)
    c = Counter()
    for line in log_text.splitlines():
        parts = line.split("\t")
        if len(parts) != 2:
            continue
        try:
            d = dt.date.fromisoformat(parts[0])
        except ValueError:
            continue
        if d >= since:
            c[parts[1].strip()] += 1
    return c


def choose(rows, current, clicks, state, cfg):
    total = cfg.get("count", 100)
    # 1. 残す：今ページにある動画のうち、クリック数が多い順に keep_top 件まで
    ranked = sorted((k for k in current if clicks[k] >= cfg.get("min_clicks", 1)),
                    key=lambda k: -clicks[k])
    kept = ranked[:min(cfg.get("keep_top", 30), total)]

    # 2. 入れ替え：Excel からランダム。直近 avoid_recent_days 日に出したものは後回し
    recent = set()
    for d, keys in state.get("history", {}).items():
        if (TODAY - dt.date.fromisoformat(d)).days < cfg.get("avoid_recent_days", 3):
            recent.update(keys)
    kept_set = set(kept)
    pool = [r for r in rows if r[0] not in kept_set]
    rng = random.Random(f"{TODAY}:{cfg.get('seed', '')}")  # 同じ日なら dry-run と本番で同じ結果
    fresh = [r for r in pool if r[0] not in recent]
    stale = [r for r in pool if r[0] in recent]
    rng.shuffle(fresh)
    rng.shuffle(stale)
    new = (fresh + stale)[: total - len(kept)]
    return kept, new


# ---------- FTP ----------

def ftp_connect(cfg):
    f = cfg["ftp"]
    cls = ftplib.FTP_TLS if f.get("protocol", "ftps") == "ftps" else ftplib.FTP
    ftp = cls()
    ftp.encoding = "utf-8"
    ftp.connect(f["host"], f.get("port", 21), timeout=60)
    ftp.login(f["user"], f["password"])
    if isinstance(ftp, ftplib.FTP_TLS):
        ftp.prot_p()
    return ftp


def ftp_download(ftp, remote, missing_ok=False):
    buf = io.BytesIO()
    try:
        ftp.retrbinary(f"RETR {remote}", buf.write)
    except ftplib.error_perm:
        if missing_ok:
            return ""
        raise
    return buf.getvalue().decode("utf-8", errors="replace")


def ftp_upload(ftp, remote, data):
    if isinstance(data, str):
        data = data.encode("utf-8")
    ftp.storbinary(f"STOR {remote}", io.BytesIO(data))


def ftp_makedirs(ftp, path):
    cur = ""
    for part in path.strip("/").split("/"):
        cur += "/" + part
        try:
            ftp.mkd(cur)
        except ftplib.error_perm:
            pass  # 既にある


def upload_server_files(ftp, root):
    """クリック計測用のファイル（server_files/ 以下）をサイト直下に同じ構成でアップロード"""
    base = HERE / "server_files"
    for f in sorted(base.rglob("*")):
        if f.is_file():
            rel = f.relative_to(base).as_posix()
            remote = posixpath.join(root, rel)
            ftp_makedirs(ftp, posixpath.dirname(remote))
            ftp_upload(ftp, remote, f.read_bytes())


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    cfg = load_config()
    state_path = HERE / "state.json"
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    if state.get("last_run") == str(TODAY) and not (args.force or args.dry_run):
        log(f"本日（{TODAY}）は実行済みのためスキップ（やり直すときは --force）")
        return

    rows = read_rows(cfg)
    titles = {k: t for k, t, _ in rows}

    remote = cfg["ftp"]["remote_html"]
    root = posixpath.dirname(remote)
    ftp = ftp_connect(cfg)
    try:
        page = ftp_download(ftp, remote)
        backups = HERE / "backups"
        backups.mkdir(exist_ok=True)
        (backups / f"index_{NOW:%Y%m%d_%H%M%S}.html").write_text(page, encoding="utf-8")
        for old in sorted(backups.glob("index_*.html"))[:-cfg.get("keep_backups", 30)]:
            old.unlink()

        clicks = Counter()
        if cfg.get("tracking", True):
            clicks = count_clicks(ftp_download(ftp, posixpath.join(root, "api/clicks.log"), missing_ok=True),
                                  cfg.get("click_window_days", 7))

        current = current_items(page, cfg)
        kept, new = choose(rows, current, clicks, state, cfg)
        if not kept and not new:
            sys.exit("掲載する動画が 0 件になるため中断します（Excel の中身を確認してください）。")

        items = [current[k] for k in kept] + [render(t, v, cfg) for _, t, v in new]
        if cfg.get("order") == "shuffle":
            random.Random(str(TODAY)).shuffle(items)
        new_page = rebuild_page(page, items, cfg)

        for k in kept:
            log(f"  残す: {k}（{clicks[k]} クリック）{titles.get(k, '')[:30]}")
        log(f"残す {len(kept)} 件 ＋ 入れ替え {len(new)} 件 ＝ {len(items)} 件"
            f"（直近{cfg.get('click_window_days', 7)}日のクリック合計 {sum(clicks.values())}）")
        if len(items) < cfg.get("count", 100):
            log(f"注意: Excel の件数が足りず {len(items)} 件になりました")

        if args.dry_run:
            (HERE / "preview.html").write_text(new_page, encoding="utf-8")
            log("[dry-run] preview.html を作成しました（アップロードはしていません）")
            return

        if cfg.get("tracking", True):
            upload_server_files(ftp, root)
        ftp_upload(ftp, remote, new_page)
        if cfg.get("local_copy"):
            Path(cfg["local_copy"]).write_text(new_page, encoding="utf-8")
    finally:
        try:
            ftp.quit()
        except Exception:
            ftp.close()

    hist = state.get("history", {})
    hist[str(TODAY)] = [k for k, _, _ in new]
    state["history"] = {d: v for d, v in sorted(hist.items())[-14:]}
    state.update(last_run=str(TODAY), last_kept=kept)
    state_path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    log(f"完了: {remote} を更新しました")


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        if e.code not in (None, 0):
            log(f"エラー: {e.code}")
        raise
    except Exception as e:
        log(f"エラー: {type(e).__name__}: {e}")
        raise
