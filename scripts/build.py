#!/usr/bin/env python3
"""
最新の X_post_sheet_YYYYMMDD.xlsx から、プラットフォーム別のページを作る。

  public/index.html            → 先頭プラットフォームのページへ転送（#v-xxx も引き継ぐ）
  public/<path>/index.html      動画ブロックの一覧（各ブロックに id="v-<ID>"）
  public/<path>/catalog.json これまで掲載したブロックの HTML（日替わりで消えた
                                 ブロックへの SNS リンクも開けるようにするため）
  public/assets/                static/ をコピー

SNS 用リンク: https://<サイト>/<path>/#v-<ID>
  <path> は PLATFORMS の "path"。商標を URL に入れないため、サービス名とは無関係な値にしている

プラットフォームごとのデータ元:
  fanza  … シート「API」。発売済みの作品から人気順位上位 KEEP_RANKED 件＋日替わりランダムで COUNT 件
  mgs    … シート「sheet」の本文にある MGS ウィジェット・mgstage.com の商品リンク ＋ シート「MGS」（あれば）
           拾った作品は public/k3wn/items.json に貯め、xlsx から消えてもページに残す
  myfans … シート「sheet」の本文にある myfans リンク ＋ シート「myfans」（あれば）
  シート「MGS」「myfans」の列: ID / タイトル / URL / 画像URL / 埋め込みHTML / 説明

使い方:
  python scripts/build.py              リポジトリ直下の最新 xlsx を使う
  python scripts/build.py path.xlsx    xlsx を指定する
"""
import datetime as dt
import html
import json
import random
import re
import shutil
import sys
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public"
JST = dt.timezone(dt.timedelta(hours=9))
TODAY = dt.datetime.now(JST).date()

SITE_TITLE = "エロ三昧タイム"
COUNT = 100
KEEP_RANKED = 30
CATALOG_MAX = 3000  # catalog.json に残すブロック数の上限（最後に掲載した日が新しい順）
# FANZA サンプル動画プレイヤー（litevideo）用のアフィリエイト ID
PLAYER_AFFI_ID = "NUBh0jCpE9kc-001"
PLAYER_URL = ("https://www.dmm.co.jp/litevideo/-/part/=/affi_id={affi}"
              "/cid={cid}/size=1280_720/")

PLATFORMS = [
    {"key": "fanza", "path": "r8tq", "name": "えろざんまい", "cta": "詳細を見る"},
    {"key": "mgs", "path": "k3wn", "name": "MGS", "cta": "MGS動画で詳細を見る"},
    {"key": "myfans", "path": "p6hz", "name": "myfans", "cta": "myfansで見る"},
]

MGS_RE = re.compile(r'<div class="[^"]*"></div><script[^>]*mgs_Widget_affiliate[^>]*></script>')
# X 投稿（fetch_mgs.py → publish_gassheet.py）の本文末尾にある MGS 商品ページのリンク
MGS_URL_RE = re.compile(r"https://www\.mgstage\.com/product/product_detail/([A-Za-z0-9_-]+)/\S*")
MGS_BACKFILL = 7  # items.json を初めて作るとき、過去の xlsx を何個さかのぼるか
MYFANS_RE = re.compile(r"https://link\.affiliate\.myfans\.jp/r/(\w+)")
URL_RE = re.compile(r"https?://\S+")


# ---------- xlsx ----------

def find_xlsx(argv):
    if len(argv) > 1:
        return Path(argv[1])
    files = sorted(ROOT.glob("X_post_sheet_*.xlsx"))
    if not files:
        sys.exit("X_post_sheet_*.xlsx が見つかりません。")
    return files[-1]


def sheet_rows(wb, name):
    """シートを {見出し: 値} の dict のリストで返す。シートがなければ空。"""
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    out = []
    for row in rows:
        d = {}
        for h, v in zip(header, row):
            if h:
                d[h] = v.strip() if isinstance(v, str) else v
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def safe_id(s):
    return re.sub(r"[^0-9A-Za-z_-]", "-", str(s)).strip("-").lower()


def item(id, title, url, **kw):
    base = {"id": safe_id(id), "title": title or "", "url": url or "", "image": "",
            "player": "", "embed": "", "desc": "", "meta": "", "rank": None}
    base.update(kw)
    return base


# ---------- プラットフォーム別の読み込み ----------

def load_fanza(wb):
    items, seen = [], set()
    for r in sheet_rows(wb, "API"):
        cid = r.get("content_id")
        released = as_date(r.get("発売日"))
        if not cid or cid in seen or released is None or released > TODAY:
            continue
        seen.add(cid)
        rank = r.get("人気順位")
        items.append(item(
            cid, r.get("タイトル"),
            r.get("アフィリエイトURL") or r.get("通常URL"),
            image=r.get("画像URL(大)") or r.get("画像URL(リスト)") or "",
            player=PLAYER_URL.format(affi=PLAYER_AFFI_ID, cid=cid),
            meta=" / ".join(str(x) for x in (r.get("女優"), released.isoformat()) if x),
            rank=int(rank) if isinstance(rank, (int, float)) else None,
        ))
    ranked = sorted((i for i in items if i["rank"]), key=lambda i: i["rank"])
    picked = ranked[:KEEP_RANKED]
    rest = [i for i in items if i not in picked]
    rng = random.Random(TODAY.isoformat())
    picked += rng.sample(rest, min(COUNT - len(picked), len(rest)))
    return picked


def post_rows(wb):
    """シート「sheet」（X 投稿）の行を新しい投稿日順に返す。"""
    rows = [r for r in sheet_rows(wb, "sheet") if isinstance(r.get("本文"), str)]
    rows.sort(key=lambda r: as_date(r.get("投稿日")) or dt.date.min, reverse=True)
    return rows


def post_texts(wb):
    """シート「sheet」（X 投稿）の本文を新しい投稿日順に返す。"""
    return [(r["本文"].replace('""', '"'), as_date(r.get("投稿日"))) for r in post_rows(wb)]


def post_media(r):
    """G 列「画像または動画ファイル（…）」の値（見出しが長いので前方一致で探す）。"""
    for k, v in r.items():
        if k.startswith("画像") and isinstance(v, str) and v.startswith("http"):
            return v
    return ""


def caption(text):
    """投稿本文から URL・タグ・ウィジェットを除いた見出しと説明を作る。"""
    text = MGS_RE.sub("", text)
    text = URL_RE.sub("", text)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return "", ""
    return lines[0], " ".join(lines[1:])


def load_extra_sheet(wb, name):
    out = []
    for r in sheet_rows(wb, name):
        id_ = r.get("ID") or r.get("URL")
        if not id_:
            continue
        out.append(item(id_, r.get("タイトル"), r.get("URL"),
                        image=r.get("画像URL") or "", embed=r.get("埋め込みHTML") or "",
                        desc=r.get("説明") or ""))
    return out


def mgs_candidates():
    """xposts/*/mgs_candidates.json（fetch_mgs.py の出力）の作品情報 {品番: {...}}。
    xposts/ は .gitignore なので、この PC にあるときだけタイトル・出演者を補う。"""
    out = {}
    for p in sorted((ROOT / "xposts").glob("*/mgs_candidates.json")):
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except ValueError:
            continue
        for i in data.get("items", []):
            if i.get("pid"):
                out[i["pid"]] = i
    return out


def load_mgs(wb):
    items = load_extra_sheet(wb, "MGS")
    cands = None
    for r in post_rows(wb):
        text = r["本文"].replace('""', '"')
        posted = as_date(r.get("投稿日"))
        day = posted.isoformat() if posted else ""
        head, desc = caption(text)
        m = MGS_RE.search(text)
        if m:
            src = re.search(r'src="([^"]+)"', m.group(0)).group(1)
            q = parse_qs(urlparse(html.unescape(src)).query)
            pid = (q.get("p") or [src])[0]
            title = (q.get("s") or [""])[0]
            items.append(item(
                pid, title or head, f"https://www.mgstage.com/product/product_detail/{pid}/",
                embed=m.group(0), desc=head if title else desc, meta=day, posted=day))
            continue
        m = MGS_URL_RE.search(text)
        if not m:
            continue
        if cands is None:
            cands = mgs_candidates()
        pid = m.group(1)
        c = cands.get(pid, {})
        title = c.get("title", "")
        items.append(item(
            pid, title or head, m.group(0),
            image=post_media(r) or c.get("image", ""),
            desc=head if title else desc,
            meta=" / ".join(x for x in (c.get("actor"), day) if x), posted=day))
    return dedupe(items)


def accumulate(path, items, first, backfill=()):
    """これまで拾った作品を items.json に貯めて、全件を「初めて載せた日」の新しい順で返す。
    {id: {カード表示の項目..., "first": "YYYY-MM-DD"}}。今回拾った作品は内容を更新し、
    拾えなかった作品も消さない。items.json がまだ無いときは backfill の
    [(日付, items), ...]（古い順）を先に入れる。"""
    try:
        store = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, ValueError):
        store = {}
        for day, old in backfill:
            for it in old:
                store[it["id"]] = {**it, "first": store.get(it["id"], {}).get("first", day)}
    for it in items:
        store[it["id"]] = {**it, "first": store.get(it["id"], {}).get("first", first)}
    path.write_text(json.dumps(store, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return sorted(store.values(), key=lambda i: (i["first"], i.get("posted", "")), reverse=True)


def xlsx_date(path):
    m = re.search(r"(\d{4})(\d{2})(\d{2})", path.name)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else TODAY.isoformat()


def mgs_backfill(current):
    """リポジトリ直下に残っている過去の xlsx（最大 MGS_BACKFILL 個）から MGS 作品を拾う。"""
    out = []
    for p in sorted(ROOT.glob("X_post_sheet_*.xlsx"))[-MGS_BACKFILL:]:
        if p.resolve() == current.resolve():
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        out.append((xlsx_date(p), load_mgs(wb)))
        wb.close()
    return out


def load_myfans(wb):
    items = load_extra_sheet(wb, "myfans")
    for text, posted in post_texts(wb):
        m = MYFANS_RE.search(text)
        if not m:
            continue
        head, desc = caption(text)
        items.append(item(m.group(1), head, m.group(0), desc=desc,
                          meta=posted.isoformat() if posted else ""))
    return dedupe(items)


def dedupe(items):
    seen, out = set(), []
    for i in items:
        if i["id"] and i["id"] not in seen:
            seen.add(i["id"])
            out.append(i)
    return out


LOADERS = {"fanza": load_fanza, "mgs": load_mgs, "myfans": load_myfans}


# ---------- HTML ----------

def render_media(it, platform):
    e = html.escape
    if it["player"]:
        badge = f'<span class="rank">人気{it["rank"]}位</span>' if it["rank"] else ""
        return (f'<button class="media" type="button" data-player="{e(it["player"])}" '
                f'aria-label="{e(it["title"])} のサンプル動画を再生">'
                f'<img src="{e(it["image"])}" alt="" loading="lazy" decoding="async">'
                f'<span class="play" aria-hidden="true"></span>{badge}</button>')
    if it["embed"]:
        # ウィジェットはページ内に 1 つしか置けないものがあるので iframe で分離する
        doc = f'<!doctype html><meta charset="utf-8"><style>body{{margin:0}}</style>{it["embed"]}'
        return (f'<iframe class="media embed" srcdoc="{e(doc)}" loading="lazy" '
                f'title="{e(it["title"])}"></iframe>')
    if it["image"]:
        return (f'<a class="media" href="{e(it["url"])}" target="_blank" rel="sponsored nofollow noopener">'
                f'<img src="{e(it["image"])}" alt="" loading="lazy" decoding="async"></a>')
    return (f'<a class="media blank" href="{e(it["url"])}" target="_blank" rel="sponsored nofollow noopener">'
            f'<span>{e(platform["name"])}</span></a>')


def render_card(it, platform):
    e = html.escape
    desc = f'<p class="desc">{e(it["desc"])}</p>' if it["desc"] else ""
    meta = f'<p class="meta">{e(it["meta"])}</p>' if it["meta"] else ""
    return f"""<li class="card" id="v-{it['id']}">
  {render_media(it, platform)}
  <div class="body">
    <p class="title">{e(it['title'])}</p>{desc}{meta}
    <div class="actions">
      <a class="cta" href="{e(it['url'])}" target="_blank" rel="sponsored nofollow noopener">{e(platform['cta'])}</a>
      <button class="copy" type="button" data-id="v-{it['id']}">リンクをコピー</button>
    </div>
  </div>
</li>"""


def render_menu(current):
    links = "\n".join(
        f'        <li><a href="../{p["path"]}/"{" aria-current=\"page\"" if p["key"] == current["key"] else ""}>{html.escape(p["name"])}</a></li>'
        for p in PLATFORMS)
    return f"""<details class="menu">
      <summary>{html.escape(current["name"])}</summary>
      <ul>
{links}
      </ul>
    </details>"""


def render_page(platform, items):
    cards = "\n".join(render_card(i, platform) for i in items)
    if not cards:
        cards = '<li class="empty">準備中です。</li>'
    template = (ROOT / "scripts" / "template.html").read_text(encoding="utf-8")
    return (template
            .replace("{{SITE_TITLE}}", html.escape(SITE_TITLE))
            .replace("{{PLATFORM}}", html.escape(platform["name"]))
            .replace("{{MENU}}", render_menu(platform))
            .replace("{{DATE}}", f"{TODAY.year}年{TODAY.month}月{TODAY.day}日")
            .replace("{{COUNT}}", str(len(items)))
            .replace("{{CARDS}}", cards))


def update_catalog(path, platform, items):
    """これまで掲載したブロックを貯める。{id: {"html": ..., "last": "YYYY-MM-DD"}}"""
    try:
        catalog = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, ValueError):
        catalog = {}
    for it in items:
        catalog[f"v-{it['id']}"] = {"html": render_card(it, platform), "last": TODAY.isoformat()}
    keep = sorted(catalog.items(), key=lambda kv: kv[1]["last"], reverse=True)[:CATALOG_MAX]
    path.write_text(json.dumps(dict(sorted(keep)), ensure_ascii=False, separators=(",", ":")),
                    encoding="utf-8")


REDIRECT = """<!doctype html>
<html lang="ja"><head><meta charset="utf-8"><meta name="robots" content="noindex">
<title>{title}</title>
<script>location.replace('{path}/' + location.hash);</script>
<meta http-equiv="refresh" content="0; url={path}/">
</head><body><a href="{path}/">動画一覧へ</a></body></html>
"""


def main():
    path = find_xlsx(sys.argv)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    OUT.mkdir(exist_ok=True)
    shutil.copytree(ROOT / "static", OUT / "assets", dirs_exist_ok=True)
    (OUT / "index.html").write_text(REDIRECT.format(title=html.escape(SITE_TITLE), path=PLATFORMS[0]["path"]), encoding="utf-8")
    for p in PLATFORMS:
        items = LOADERS[p["key"]](wb)
        d = OUT / p["path"]
        d.mkdir(exist_ok=True)
        shown = items
        if p["key"] == "mgs":
            # MGS は一度載せた作品を残し続ける（public/k3wn/items.json）
            store = d / "items.json"
            shown = accumulate(store, items, TODAY.isoformat(),
                               backfill=() if store.exists() else mgs_backfill(path))
        (d / "index.html").write_text(render_page(p, shown), encoding="utf-8")
        update_catalog(d / "catalog.json", p, items)
        # ログ（cp932）で化けないよう ASCII で出す
        extra = f" (from xlsx {len(items)})" if shown is not items else ""
        print(f"{path.name}: {p['key']} {len(shown)} items{extra}")


if __name__ == "__main__":
    main()
